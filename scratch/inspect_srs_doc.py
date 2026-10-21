import docx

path = r'C:\Users\Abyureddy\OneDrive\Desktop\SDC2\SDC Review 1\SRS_V1_ (2).docx'
doc = docx.Document(path)

print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

headers = []
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if p.style.name.startswith('Heading') or (len(txt) < 80 and any(txt.startswith(x) for x in ['1.', '2.', '3.', '4.', '5.', '6.', 'APPENDIX', 'TABLE'])):
        headers.append((i, p.style.name, txt))

print(f"Found {len(headers)} section headers:")
for idx, style, h in headers[:40]:
    print(f"  [{idx}] ({style}): {h}")

for t_idx, table in enumerate(doc.tables):
    print(f"\n--- Table {t_idx+1}: {len(table.rows)} rows x {len(table.columns)} cols ---")
    for r in table.rows[:3]:
        row_txt = [c.text.strip().replace('\n', ' ') for c in r.cells]
        print(f"   {row_txt[:4]}")
