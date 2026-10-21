import pptx

path = r'C:\Users\Abyureddy\OneDrive\Desktop\SDC2\SDC Review 1\malla nene chesaa guysss.pptx'
prs = pptx.Presentation(path)
print(f'Slide count: {len(prs.slides)}')
for i, slide in enumerate(prs.slides):
    print(f'\n=== SLIDE {i+1} ===')
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                t = p.text.strip()
                if t:
                    safe_t = t.encode('ascii', 'replace').decode('ascii')
                    print(f'  [{shape.name} - ID {shape.shape_id}]: {safe_t}')
