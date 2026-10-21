import pptx

prs = pptx.Presentation(r'C:\Users\Abyureddy\OneDrive\Desktop\SDC2\SDC Review 1\PPT.pptx')
for idx in [1, 2, 3, 4, 5, 6]:
    s = prs.slides[idx]
    print(f"=== SLIDE {idx+1} ===")
    for shape in s.shapes:
        print(f"Shape {shape.name}: left={shape.left.inches:.2f}, top={shape.top.inches:.2f}, w={shape.width.inches:.2f}, h={shape.height.inches:.2f}")
        if shape.has_text_frame:
            print(f"  Text: {shape.text_frame.text[:80]}")
