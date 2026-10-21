import docx

path = r"C:\Users\Abyureddy\OneDrive\Desktop\SDC2\SDC Review 1\SRS_V1_ (2).docx"
doc = docx.Document(path)

for p in doc.paragraphs:
    for r in p.runs:
        if chr(8211) in r.text or chr(65533) in r.text or "\x00" in r.text:
            cleaned = r.text.replace(chr(8211), "").replace(chr(65533), "").replace("\x00", "")
            r.text = cleaned

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    if chr(8211) in r.text or chr(65533) in r.text or "\x00" in r.text:
                        cleaned = r.text.replace(chr(8211), "").replace(chr(65533), "").replace("\x00", "")
                        r.text = cleaned

doc.save(path)
print("En-dash interleaved characters stripped successfully.")
