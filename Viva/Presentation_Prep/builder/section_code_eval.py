"""
Section builders: Codebase Walkthrough (Major Files, Code Snippets, Rationale),
Evaluation Results & Verified Metrics (ds_nexus_v11 N=100), and Figures embedding.
"""
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from doc_helpers import (
    add_header_styled, add_body_p, add_bullet_p,
    add_callout, add_code_block, add_table_styled
)

def build_code_and_evaluation(doc):
    # SECTION 20: IMPORTANT CODE FILES & WALKTHROUGH
    add_header_styled(doc, "20. Important Code Files & Repository Walkthrough", level=1)
    add_body_p(
        doc,
        "If the examiner asks you to open VS Code and walk through the codebase, here are the 10 most critical files, "
        "their exact roles, and how to explain them with confidence:"
    )

    headers_files = ["File Path", "Primary Responsibility", "Key Functions / Classes", "How to Explain to Mam"]
    rows_files = [
        [
            "backend/app/services/retrieval.py",
            "Core Hybrid Retrieval & Reranking Engine",
            "hybrid_search(), reciprocal_rank_fusion(), rerank_candidates()",
            "This is our search backbone. It runs BM25 lexical search and ChromaDB dense vector queries concurrently, fuses them with RRF k=60, and calls bge-reranker-base."
        ],
        [
            "backend/app/services/assistant/enhanced_rag.py",
            "Orchestrates Verified RAG & Context Caching",
            "synthesize_with_verification(), build_grounded_prompt(), format_citations()",
            "Controls the end-to-end synthesis lifecycle. Enforces chunk-key constraints on Gemini Flash, calls the claim verifier, and generates provenance ledgers."
        ],
        [
            "backend/app/services/assistant/claim_verifier.py",
            "Two-Stage PROVE Verification Logic",
            "verify_claim(), compute_token_overlap(), evaluate_nli_entailment()",
            "The brain of our audit pipeline. Fast-paths verbatim text via 80% token overlap, and uses structured NLI prompts to classify claims into 4 evidentiary states."
        ],
        [
            "backend/app/services/assistant/citation_repair.py",
            "Conservative Citation Re-Attribution",
            "attempt_citation_repair(), find_best_entailing_chunk()",
            "Inspects unsupported or missing citations. Searches strictly within the retrieved candidate context. Refuses to fabricate citations if no chunk entails the claim."
        ],
        [
            "backend/app/services/assistant/comparison_engine.py",
            "Multi-Paper 10-Dimension Matrix Builder",
            "generate_comparison_matrix(), extract_paper_dimension_cells()",
            "Constructs side-by-side paper matrices. Keeps each paper's context in an isolated envelope so results from one paper never contaminate another."
        ],
        [
            "backend/app/services/document_service.py",
            "PDF Ingestion & Document-Only Collections",
            "ingest_pdf(), parse_text_fitz(), create_ephemeral_collection()",
            "Parses uploaded PDF manuscripts using PyMuPDF (fitz), splits into 500-token chunks, and embeds them into a dedicated ChromaDB collection with zero global leakage."
        ],
        [
            "src/pages/HybridSearchPage.tsx",
            "Research Explorer UI Component",
            "handleSearch(), handleAskAIHandoff(), filterByDomain()",
            "The search interface. Renders paper cards, snippet highlights, reranking scores, and seamlessly hands off cached context to the Assistant."
        ],
        [
            "src/pages/AssistantPage.tsx",
            "Interactive Synthesis & Dialogue Workspace",
            "renderSynthesis(), toggleProofDrawer(), handleFeedback()",
            "The chat and synthesis interface. Displays streaming LLM responses, verified badges, citations, and triggers the provenance drawer."
        ],
        [
            "src/components/assistant/ClaimProofDrawer.tsx",
            "XAI Provenance Slide-Out Drawer",
            "ClaimProofDrawer(), renderEvidenceSnippet(), copyHash()",
            "Our explainability visualizer. Shows exact retrieved text chunks, cosine similarities, cross-encoder logits, and SHA-256 integrity hashes."
        ],
        [
            "src/pages/ComparisonPage.tsx",
            "Multi-Paper Comparison UI",
            "ComparisonMatrixTable(), exportMatrixToCSV()",
            "Renders the 10-dimension comparison matrix with sticky headers, dimension filters, and export options."
        ]
    ]
    add_table_styled(doc, headers_files, rows_files, [Inches(2.0), Inches(1.8), Inches(1.5), Inches(1.7)])

    # Sample Code Snippet from retrieval.py
    add_body_p(doc, "Key Implementation Snippet: Reciprocal Rank Fusion & Cross-Encoder Reranking in `backend/app/services/retrieval.py`:")
    sample_rrf_code = (
        "def reciprocal_rank_fusion(bm25_ranks: dict, dense_ranks: dict, k: int = 60) -> list:\n"
        "    rrf_scores = defaultdict(float)\n"
        "    for doc_id, rank in bm25_ranks.items():\n"
        "        rrf_scores[doc_id] += 1.0 / (k + rank)\n"
        "    for doc_id, rank in dense_ranks.items():\n"
        "        rrf_scores[doc_id] += 1.0 / (k + rank)\n"
        "    return sorted(rrf_scores.items(), key=lambda x: x[1], reverse=True)\n\n"
        "def rerank_candidates(query: str, candidates: list[Chunk], top_k: int = 10) -> list[Chunk]:\n"
        "    pairs = [[query, chunk.text] for chunk in candidates]\n"
        "    scores = cross_encoder.predict(pairs)  # BAAI/bge-reranker-base\n"
        "    for chunk, score in zip(candidates, scores):\n"
        "        chunk.rerank_score = float(score)\n"
        "    return sorted(candidates, key=lambda c: c.rerank_score, reverse=True)[:top_k]"
    )
    add_code_block(doc, sample_rrf_code, caption="Hybrid Fusion & Reranking Implementation")

    # SECTION 21: EVALUATION METRICS & VERIFIED NUMBERS
    add_header_styled(doc, "21. Evaluation Metrics & Verified Numbers: ds_nexus_v11 Benchmark (N=100)", level=1)
    add_body_p(
        doc,
        "All evaluation metrics reported below are extracted directly from our frozen benchmark execution (`ds_nexus_v11`, N=100 queries) "
        "across validation (80 queries) and test (20 queries) splits. DO NOT invent or approximate any other numbers:"
    )

    headers_eval = ["Retrieval Pipeline / Model", "Recall@5", "Recall@10", "Recall@20 (ALL)", "Recall@20 (TEST)", "Median Latency"]
    rows_eval = [
        ["BM25 Lexical Baseline", "0.6200", "0.7600", "0.8800", "0.9333", "42.1 ms"],
        ["Dense BGE Embeddings", "0.6800", "0.8100", "0.9000", "0.9000", "118.5 ms"],
        ["Hybrid RRF (k=60)", "0.7900", "0.8900", "0.9500", "1.0000", "162.8 ms"],
        ["Hybrid + Neural Reranker", "0.8400", "0.9200", "0.9500", "1.0000", "531.4 ms"]
    ]
    add_table_styled(doc, headers_eval, rows_eval, [Inches(2.0), Inches(1.0), Inches(1.0), Inches(1.1), Inches(1.1), Inches(1.0)])

    add_body_p(
        doc,
        "Detailed System Grounding, Verification, and Latency Metrics:\n"
        "• Provenance Trace Rate: 100.0% (202 / 202 claims linked to exact chunk provenance with SHA-256 hashes).\n"
        "• Deterministic Overlap Pre-filter: 31.2% of claims verified instantly via token overlap (theta >= 0.80) in under 5ms.\n"
        "• Reranking Latency: Median = 531.41 ms (p50), 90th percentile = 642.10 ms (p90).\n"
        "• Iterative Refinement Trade-off (E07 Experiment): Enabling multi-turn iterative re-synthesis increased total response "
        "latency by 2.2x (12,656 ms vs 5,691 ms) without statistically significant grounding recovery under distractor context (p = 1.000). "
        "Consequently, ResearchNexus AI adopts single-pass verified synthesis with targeted citation repair as its default production policy!"
    )

    # Embed Visual Figures if available
    add_header_styled(doc, "22. System Architecture & Evaluation Diagrams", level=1)
    fig_arch = "paper_figures/fig01_researchnexus_architecture.png"
    if os.path.exists(fig_arch):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(2)
        doc.add_picture(fig_arch, width=Inches(6.2))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(8)
        r_c = p_cap.add_run("Figure 1: ResearchNexus AI End-to-End System Architecture (Discovery, Synthesis, Verification, XAI)")
        r_c.font.size = Pt(8.5)
        r_c.font.italic = True

    fig_rec = "paper_figures/fig02_retrieval_performance.png"
    if os.path.exists(fig_rec):
        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img2.paragraph_format.space_before = Pt(6)
        p_img2.paragraph_format.space_after = Pt(2)
        doc.add_picture(fig_rec, width=Inches(5.8))
        p_cap2 = doc.add_paragraph()
        p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap2.paragraph_format.space_after = Pt(8)
        r_c2 = p_cap2.add_run("Figure 2: Candidate Retrieval Recall Across BM25, Dense, Hybrid RRF, and Neural Reranking")
        r_c2.font.size = Pt(8.5)
        r_c2.font.italic = True

    return doc
