"""
Document builder helper functions for formatting, tables, callouts, and styling.
"""
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def setup_page(doc):
    for sec in doc.sections:
        sec.top_margin = Inches(0.75)
        sec.bottom_margin = Inches(0.75)
        sec.left_margin = Inches(0.75)
        sec.right_margin = Inches(0.75)

def add_callout(doc, title, text, bg_hex="EBF5FB", border_hex="2980B9"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.8)
    cell = tbl.cell(0, 0)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    cell._tc.get_or_add_tcPr().append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p_run = p.add_run(f"★ {title}\n")
    p_run.bold = True
    p_run.font.name = 'Calibri'
    p_run.font.size = Pt(10.5)
    p_run.font.color.rgb = RGBColor.from_string(border_hex)
    
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(4)
    p2_run = p2.add_run(text)
    p2_run.font.name = 'Calibri'
    p2_run.font.size = Pt(10)
    p2_run.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
    
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_str, caption=None):
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(4)
        p_cap.paragraph_format.space_after = Pt(2)
        r_cap = p_cap.add_run(f"Code Listing: {caption}")
        r_cap.bold = True
        r_cap.font.name = 'Calibri'
        r_cap.font.size = Pt(9.5)
        r_cap.font.color.rgb = RGBColor(0x1B, 0x4F, 0x72)

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.8)
    cell = tbl.cell(0, 0)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F4F6F7"/>')
    cell._tc.get_or_add_tcPr().append(shd)
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="0" w:color="7F8C8D"/><w:top w:val="single" w:sz="4" w:space="0" w:color="BDC3C7"/><w:right w:val="single" w:sz="4" w:space="0" w:color="BDC3C7"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="BDC3C7"/></w:tcBorders>')
    cell._tc.get_or_add_tcPr().append(tcBorders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(code_str.strip())
    r.font.name = 'Consolas'
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0x21, 0x2F, 0x3D)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(4)

def add_header_styled(doc, text, level=1):
    h = doc.add_heading(level=level)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    r.font.name = 'Calibri'
    if level == 1:
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        r.font.size = Pt(16)
        r.bold = True
        r.font.color.rgb = RGBColor(0x0E, 0x4D, 0x92) # Navy
    elif level == 2:
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        r.font.size = Pt(13)
        r.bold = True
        r.font.color.rgb = RGBColor(0x1B, 0x4F, 0x72) # Steel Blue
    elif level == 3:
        h.paragraph_format.space_before = Pt(6)
        h.paragraph_format.space_after = Pt(2)
        r.font.size = Pt(11)
        r.bold = True
        r.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50) # Charcoal
    return h

def add_body_p(doc, text, bold_prefix=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(10)
        r_pre.font.color.rgb = RGBColor(0x1A, 0x25, 0x2C)
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
    return p

def add_bullet_p(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(10)
        r_pre.font.color.rgb = RGBColor(0x1A, 0x25, 0x2C)
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
    return p

def add_table_styled(doc, headers, rows_data, col_widths=None):
    tbl = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    # Format header row
    hdr_cells = tbl.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].paragraphs[0].text = ""
        hdr_cells[i].paragraphs[0].paragraph_format.space_before = Pt(3)
        hdr_cells[i].paragraphs[0].paragraph_format.space_after = Pt(3)
        r = hdr_cells[i].paragraphs[0].add_run(h_text)
        r.bold = True
        r.font.name = 'Calibri'
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1B4F72"/>')
        hdr_cells[i]._tc.get_or_add_tcPr().append(shd)

    # Format data rows
    for r_idx, row in enumerate(rows_data):
        row_cells = tbl.rows[r_idx + 1].cells
        bg_color = "F9EBEA" if r_idx % 2 == 1 else "FFFFFF"
        # subtle alternate
        bg_fill = "F4F6F7" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].paragraphs[0].text = ""
            row_cells[c_idx].paragraphs[0].paragraph_format.space_before = Pt(2)
            row_cells[c_idx].paragraphs[0].paragraph_format.space_after = Pt(2)
            r = row_cells[c_idx].paragraphs[0].add_run(str(val))
            r.font.name = 'Calibri'
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_fill}"/>')
            row_cells[c_idx]._tc.get_or_add_tcPr().append(shd)

    # Set column widths if supplied
    if col_widths:
        for idx, width in enumerate(col_widths):
            w_val = Inches(width) if isinstance(width, (int, float)) and not hasattr(width, 'emu') else width
            tbl.columns[idx].width = w_val
            for cell in tbl.columns[idx].cells:
                cell.width = w_val

    # Add light borders
    for row in tbl.rows:
        for cell in row.cells:
            tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="4" w:space="0" w:color="D5D8DC"/><w:top w:val="single" w:sz="4" w:space="0" w:color="D5D8DC"/><w:right w:val="single" w:sz="4" w:space="0" w:color="D5D8DC"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="D5D8DC"/></w:tcBorders>')
            cell._tc.get_or_add_tcPr().append(tcBorders)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(4)
    return tbl
