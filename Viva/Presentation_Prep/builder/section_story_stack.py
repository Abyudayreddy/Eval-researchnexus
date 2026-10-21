"""
Section builders: Cover, Project Story, Problem Formulation, Architecture & Flow,
Module Breakdown, and Retrieval Engine Deep Dive.
"""
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from doc_helpers import (
    add_header_styled, add_body_p, add_bullet_p,
    add_callout, add_code_block, add_table_styled
)

def build_cover_and_intro(doc):
    # Title & Header
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(36)
    title_p.paragraph_format.space_after = Pt(8)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = title_p.add_run("RESEARCHNEXUS AI\nPROJECT PRESENTATION PREPARATION & COMPLETE STUDY GUIDE")
    r_title.bold = True
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(24)
    r_title.font.color.rgb = RGBColor(0x0E, 0x4D, 0x92)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(16)
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = sub_p.add_run("Understanding the Problem, System, Architecture, Implementation, Live Demo, Evaluation and Future Direction")
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(13)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(0x56, 0x65, 0x73)

    badge_p = doc.add_paragraph()
    badge_p.paragraph_format.space_before = Pt(8)
    badge_p.paragraph_format.space_after = Pt(24)
    badge_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_badge = badge_p.add_run("Prepared for Project Presentation – Tuesday\nComprehensive Defense, Technical Architecture & Live Demo Handbook")
    r_badge.bold = True
    r_badge.font.name = 'Calibri'
    r_badge.font.size = Pt(11)
    r_badge.font.color.rgb = RGBColor(0x1B, 0x4F, 0x72)

    # Note / Purpose box
    add_callout(
        doc,
        "HOW TO USE THIS PREPARATION GUIDE",
        "This handbook is NOT an abstract paper or boilerplate report. It is your conversational, deep-dive project companion. "
        "Read it front-to-back before Tuesday: it walks through the exact genesis of ResearchNexus AI, explains every module "
        "with concrete analogies and production details, provides verbatim scripts for what to show and say during your live demo, "
        "and supplies answers to 40+ examiner questions so you can defend every architectural choice with authority.",
        bg_hex="EAF2F8",
        border_hex="2471A3"
    )

    doc.add_page_break()

    # SECTION 1: THE STORY OF THE PROJECT
    add_header_styled(doc, "1. The Story of the Project: Why and How We Built ResearchNexus AI", level=1)
    add_body_p(
        doc,
        "Every academic researcher, graduate student, and R&D engineer faces an exhausting, fractured workflow when "
        "starting a new investigation. Typically, your browser contains thirty to forty open tabs across arXiv, Google Scholar, "
        "IEEE Xplore, and PubMed. You skim abstracts, try to piece together contradictory experimental baselines, and copy-paste "
        "snippets into personal notes or spreadsheets. When generative AI arrived, researchers hoped ChatGPT or Claude could "
        "summarize this literature. But general-purpose LLMs immediately hit a critical wall: hallucinations, fabricated citations, "
        "and shallow, generic summaries that miss the nuance of experimental setups."
    )
    add_body_p(
        doc,
        "We asked a fundamental question: Can we build a unified research intelligence platform that feels like an experienced "
        "co-investigator? An AI platform that does not just generate plausible text, but anchors every single claim directly into peer-reviewed "
        "literature with mathematical verification, cryptographic evidence hashing, and transparent provenance."
    )
    add_body_p(
        doc,
        "That is how ResearchNexus AI was born. We built it around five core pillars:\n"
        "1. Unified Research Discovery: Moving beyond brittle keyword queries with hybrid sparse-dense retrieval.\n"
        "2. Evidence-Grounded Scientific Synthesis: A dual-stage verification engine (token-overlap check + natural language inference) "
        "that audits every generated sentence before the researcher sees it.\n"
        "3. Explainable AI (XAI) & Auditability: Interactive provenance ledgers displaying exact chunk IDs, cosine similarities, "
        "and cross-encoder reranking scores.\n"
        "4. Multi-Paper Matrix Comparison: Side-by-side analysis across 10 structured scientific dimensions (methodology, datasets, metrics, compute limits) "
        "with zero cross-paper context contamination.\n"
        "5. Self-Contained Document-Only Intelligence: When analyzing an unpublished PDF manuscript, the system enforces a strict isolated "
        "vector collection with zero fallback leakage to the global database."
    )

    add_callout(
        doc,
        "THE PRESENTATION NARRATIVE HOOK",
        "When Mam asks: 'What is the core idea of your project?', start with this narrative:\n"
        "'Mam, general-purpose LLMs produce fluent literature reviews, but up to 30% of their scientific citations are fabricated or misattributed. "
        "In academic research, an ungrounded claim can invalidate months of work. ResearchNexus AI solves this by introducing a closed-loop "
        "Verified RAG architecture. We combine hybrid sparse-dense search with cross-encoder neural reranking, extract atomic claims, "
        "audit them through a two-stage verification pipeline, and display full provenance in an interactive user workspace.'",
        bg_hex="FEF9E7",
        border_hex="F39C12"
    )

    # SECTION 2: THE PROBLEM WE ARE SOLVING
    add_header_styled(doc, "2. The Real-World Problem We Are Solving", level=1)
    add_body_p(
        doc,
        "Academic literature exploration suffers from four structural bottlenecks that existing software tools fail to resolve simultaneously:"
    )

    headers_prob = ["Bottleneck", "Traditional Limitation", "How ResearchNexus AI Solves It"]
    rows_prob = [
        [
            "Literature Fragmentation",
            "Researchers manually jump between arXiv, IEEE, Scholar, Mendeley, and Notion, losing context between search and synthesis.",
            "Unified workspace combining discovery, synthesis, comparative matrices, personal library, and document-only analysis."
        ],
        [
            "Keyword Search Blind Spots",
            "BM25 lexical search misses conceptual synonyms; pure vector search misses exact technical jargon, acronyms, and dataset names.",
            "Hybrid Search fusing BM25 and 768-dim BGE embeddings via Reciprocal Rank Fusion (RRF k=60), followed by Cross-Encoder neural reranking."
        ],
        [
            "Generative Hallucinations",
            "Standard RAG models append retrieved context to prompts, but LLMs frequently generate unsupported extrapolate claims or blend papers.",
            "Two-stage Claim Verification: Deterministic token overlap (theta=0.80) followed by semantic NLI classification into 4 strict evidentiary states."
        ],
        [
            "Opaque Citation Attribution",
            "Inline citations like [1] often point to papers that mention the topic tangentially but never state the asserted finding.",
            "Audit Ledger & Proof Drawer: Every claim maps to exact chunk offsets, chunk SHA-256 hashes, retrieval ranks, and reranking delta scores."
        ]
    ]
    add_table_styled(doc, headers_prob, rows_prob, [Inches(1.8), Inches(2.5), Inches(2.5)])

    # SECTION 3: WHAT IS RESEARCHNEXUS AI?
    add_header_styled(doc, "3. What is ResearchNexus AI? System Overview & Actual Tech Stack", level=1)
    add_body_p(
        doc,
        "ResearchNexus AI is an enterprise-grade academic research intelligence platform. It is built as a responsive, modern web application "
        "backed by high-performance asynchronous microservices and local neural models."
    )

    headers_stack = ["Layer / Subsystem", "Technology / Framework", "Specific Version & Role in Implementation"]
    rows_stack = [
        ["Frontend UI/UX", "React 18 + TypeScript + Vite", "High-responsiveness single-page application with modular atomic workspaces."],
        ["Styling & Icons", "Tailwind CSS + Lucide Icons", "Modern, dark/light theme supporting academic dashboards and interactive drawers."],
        ["Backend API Server", "FastAPI (Python 3.11+)", "High-throughput asynchronous ASGI web server with strict Pydantic schemas."],
        ["Dense Vector Embeddings", "BAAI/bge-base-en-v1.5", "768-dimensional dense vector representations optimized for retrieval tasks."],
        ["Lexical Search Engine", "rank-bm25 (BM25Okapi)", "Term frequency / inverse document frequency scoring over tokenized corpora."],
        ["Hybrid Fusion", "Reciprocal Rank Fusion (RRF)", "Rank aggregation with constant k=60 balancing lexical and semantic rank positions."],
        ["Neural Reranker", "BAAI/bge-reranker-base", "Cross-encoder scoring full (query, document) token interactions (median latency: 531ms)."],
        ["Vector Database", "ChromaDB (Persistent Client)", "Stores chunk embeddings, metadata, domain tags, and provides cosine ANN search."],
        ["Relational Metadata", "SQLite (Local) / PostgreSQL", "Stores paper catalog, user libraries, query histories, and evaluation logs."],
        ["Document Parsing", "PyMuPDF (fitz) + LlamaIndex", "Extracts clean text streams, parses section hierarchies, and creates token-bounded chunks."],
        ["Generative AI Backbone", "Google Gemini 2.5 Flash", "High-speed reasoning model with temperature=0.2 for reproducible factual synthesis."]
    ]
    add_table_styled(doc, headers_stack, rows_stack, [Inches(1.8), Inches(2.2), Inches(2.8)])

    # SECTION 4: THE PROJECT IN ONE SIMPLE FLOW
    add_header_styled(doc, "4. The Project in One Simple Flow: End-to-End Execution Pipeline", level=1)
    add_body_p(
        doc,
        "To understand ResearchNexus AI intuitively, visualize the complete lifecycle of a user query from keyboard to verified synthesis:"
    )

    flow_ascii = (
        "USER QUERY: 'What are the main causes of concept drift in continual learning?'\n"
        "   │\n"
        "   ├───► 1. HYBRID RETRIEVAL\n"
        "   │     ├─ BM25 Lexical Search (extracts exact keywords: 'concept drift', 'continual learning')\n"
        "   │     └─ BGE-Base-en-v1.5 Dense Search (768-dim cosine similarity on ChromaDB)\n"
        "   │\n"
        "   ├───► 2. RECIPROCAL RANK FUSION (RRF k=60)\n"
        "   │     └─ Merges top-N ranked lists into a single candidate pool (achieves 100% recall on TEST benchmark)\n"
        "   │\n"
        "   ├───► 3. CROSS-ENCODER RERANKING (bge-reranker-base)\n"
        "   │     └─ Scores joint (query, chunk) token cross-attention; filters out superficial keyword matches\n"
        "   │\n"
        "   ├───► 4. CACHE HANDOFF (Explorer -> Assistant Bridge)\n"
        "   │     └─ Clicking 'Ask AI' passes the retrieved chunk IDs directly; NO redundant retrieval needed!\n"
        "   │\n"
        "   ├───► 5. GENERATIVE SYNTHESIS (Gemini 2.5 Flash, temp=0.2)\n"
        "   │     └─ Drafts a comprehensive, structured scientific synthesis citing chunk keys ([chunk-xx])\n"
        "   │\n"
        "   ├───► 6. TWO-STAGE CLAIM VERIFICATION (ClaimVerifier)\n"
        "   │     ├─ Stage A: Deterministic token overlap (threshold = 0.80)\n"
        "   │     └─ Stage B: Semantic NLI (Supported / Partially Supported / Unsupported / Contradicted)\n"
        "   │\n"
        "   ├───► 7. CITATION INTEGRITY & AUDIT LEDGER\n"
        "   │     └─ Replaces weak citations conservatively; attaches cryptographic SHA-256 hashes to every claim\n"
        "   │\n"
        "   └───► 8. INTERACTIVE PROVENANCE UI (Proof Drawer)\n"
        "         └─ Displays verified badges, evidence highlights, rerank scores, and paper PDF deep-links."
    )
    add_code_block(doc, flow_ascii, caption="ResearchNexus AI Core Execution Pipeline")

    return doc
